"""Builds the Google Sheet workbook (.xlsx) for question updates.

One workbook serves several apps: a `使い方` guide tab, one shared
`manifest` tab, and one tab per question pack. Each app only reads the
packs it knows, so KIDS and TEEN can share the sheet.

    # 1) export each app's packs to CSV (same format the apps read back)
    dart run tool/sheets/export_sheets.dart ../ark_busters_kids/assets/data ../ark_busters_kids/tools/sheets
    dart run tool/sheets/export_sheets.dart ../ark_busters_teen/assets/data ../ark_busters_teen/tools/sheets
    # 2) build the workbook
    python tool/sheets/make_workbook.py out.xlsx KIDS=../ark_busters_kids TEEN=../ark_busters_teen

Then in Google Sheets: File → Import → Upload → "Replace spreadsheet".
Every cell is plain text so Google never turns 1.5 / -8 / 3/4 into numbers
or dates. Row 2 of each tab explains the columns; its `_path` cell is
empty, so the apps ignore it.

Requires: openpyxl
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TEXT = '@'
SPARE_ROWS = 200  # pre-formatted empty rows for new questions

KIDS_TITLES = {
    'vocab': 'ボキャブラリー',
    'listening': 'リスニング',
    'talk': 'トーク',
    'sentence': 'センテンス',
}

# Column meaning (exact name, or regex on the name) → Japanese description.
COLUMN_HELP = [
    (r'^id$', '番号（ほかと重ならない値）'),
    (r'^w$', '英単語'),
    (r'^lv$', '級（5・4・3）'),
    (r'^img$', '絵のファイル（空＝絵なし）'),
    (r'^voice$', '音声ファイル（空＝読み上げ）'),
    (r'^t$', '問題の種類（変えない）'),
    (r'^en$', '英語'),
    (r'^ja$', '日本語'),
    (r'^tip$', 'ポイント（解説）'),
    (r'^ans$', '正解'),
    (r'^opts\.(\d+)$', 'ほかの選択肢 {n}'),
    (r'^q$', '問題文'),
    (r'^a$', '正解'),
    (r'^choices\.(\d+)$', '選択肢 {n}（正解も入れる）'),
    (r'^e$', '解説'),
    (r'^ja\.(\d+)\.t$', '日本語のかたまり {n}'),
    (r'^ja\.(\d+)\.pick$', 'かたまり {n} に色（true / false）'),
    (r'^answer\.(\d+)\.w$', '英語 {n}（正しい順）'),
    (r'^answer\.(\d+)\.label$', '英語 {n} のラベル（誰が 等）'),
    (r'^answer\.(\d+)\.c$', '英語 {n} の色番号'),
]


def describe(col):
    for pat, text in COLUMN_HELP:
        m = re.match(pat, col)
        if m:
            n = int(m.group(1)) + 1 if m.groups() else ''
            return text.format(n=n)
    return col


def describe_path(path):
    """`grades.g5.levels.0` → `5級 レベル0`, `levels.Lv2` → `レベル2`."""
    parts = []
    for seg in path.split('.'):
        if re.fullmatch(r'g\d', seg):
            parts.append(f'{seg[1]}級')
        elif re.fullmatch(r'Lv\d', seg):
            parts.append(f'レベル{seg[2]}')
        elif seg.isdigit():
            parts.append(f'レベル{seg}')
        elif seg in ('items',):
            parts.append('全問題')
    return ' '.join(parts) or path


HEAD_FILL = PatternFill('solid', fgColor='1F3864')
HEAD_FONT = Font(bold=True, color='FFFFFF')
HELP_FILL = PatternFill('solid', fgColor='EEEEEE')
HELP_FONT = Font(italic=True, color='666666')
MANI_FILL = PatternFill('solid', fgColor='FFF2CC')


def text_cell(ws, row, col, value, **style):
    c = ws.cell(row=row, column=col, value='' if value is None else str(value))
    c.number_format = TEXT
    for k, v in style.items():
        setattr(c, k, v)
    return c


def size_columns(ws, rows, cap=42):
    widths = {}
    for r in rows:
        for i, v in enumerate(r, 1):
            w = sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(v))
            widths[i] = max(widths.get(i, 6), min(cap, w + 2))
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = w


def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.reader(f))


def main():
    out = Path(sys.argv[1])
    apps = [a.split('=', 1) for a in sys.argv[2:]]
    wb = Workbook()
    guide = wb.active
    guide.title = '使い方'
    manifest = wb.create_sheet('manifest')
    mani_head = ['pack', 'version', 'app', 'タイトル', '問題数', 'gid', 'minAppBuild', 'メモ']
    mani_rows = [mani_head]
    path_notes = []

    for app, root in apps:
        root = Path(root)
        sheets = root / 'tools' / 'sheets'
        for csv_path in sorted(sheets.glob('*.csv')):
            pack = csv_path.stem
            if pack == 'manifest':
                continue
            rows = read_csv(csv_path)
            header, data = rows[0], rows[1:]
            meta = json.loads((root / 'assets' / 'data' / f'{pack}.json').read_text(encoding='utf-8'))
            title = KIDS_TITLES.get(pack) or meta.get('title') or pack
            mani_rows.append([pack, meta.get('version', 1), app, title, len(data), '', '', ''])

            ws = wb.create_sheet(pack)
            for c, h in enumerate(header, 1):
                text_cell(ws, 1, c, h, font=HEAD_FONT, fill=HEAD_FILL)
            # Row 2: explanation (empty _path → ignored by the apps).
            for c, h in enumerate(header, 1):
                text_cell(ws, 2, c, '' if h == '_path' else describe(h),
                          font=HELP_FONT, fill=HELP_FILL)
            for r, row in enumerate(data, 3):
                for c, v in enumerate(row, 1):
                    text_cell(ws, r, c, v, alignment=Alignment(vertical='top'))
            for r in range(len(data) + 3, len(data) + 3 + SPARE_ROWS):
                for c in range(1, len(header) + 1):
                    ws.cell(row=r, column=c).number_format = TEXT
            ws.freeze_panes = 'B3'
            size_columns(ws, [header] + data[:60])
            paths = list(dict.fromkeys(row[0] for row in data))
            path_notes.append((app, pack, title, paths))

    # pricing tab: free/paid settings per app (read by ark_core pricing_sheet).
    pricing = wb.create_sheet('pricing')
    pricing_rows = [
        ['app', 'item', 'value', 'メモ'],
        ['', '', '', 'app = kids / teen。value が空の行は無視。変更後はアプリが次回起動時に取り込みます（manifest の version は不要）'],
    ]
    for app, _ in apps:
        a = app.lower()
        pricing_rows += [
            [a, 'enabled', 'false', '課金を有効にする（true / false）。false のあいだは全部無料'],
            [a, 'defaultFreeCount', '1', 'ゲーム（ステージ）ごとに無料で遊べるモード（レベル）の数。0 = 全部有料'],
            [a, 'freeCounts.<pack>', '', '特定タブだけ別の数にしたいとき（例：item を freeCounts.vocab、value を 2）'],
            [a, 'freePacks', '', '全部無料にするタブ名（カンマ区切り）'],
            [a, 'lockedPacks', '', '全部有料にするタブ名（カンマ区切り）'],
        ]
    for r, row in enumerate(pricing_rows, 1):
        for c, v in enumerate(row, 1):
            style = {'font': HEAD_FONT, 'fill': HEAD_FILL} if r == 1 else (
                {'font': HELP_FONT, 'fill': HELP_FILL} if r == 2 else ({'fill': MANI_FILL} if c == 3 else {}))
            text_cell(pricing, r, c, v, **style)
    pricing.freeze_panes = 'A3'
    size_columns(pricing, pricing_rows, cap=70)

    for r, row in enumerate(mani_rows, 1):
        for c, v in enumerate(row, 1):
            style = {'font': HEAD_FONT, 'fill': HEAD_FILL} if r == 1 else (
                {'fill': MANI_FILL} if c == 2 else {})
            text_cell(manifest, r, c, v, **style)
    manifest.freeze_panes = 'A2'
    size_columns(manifest, mani_rows)

    lines = [
        ('ARK BUSTERS 問題データ', 'bold'),
        ('', None),
        ('■ 問題を直す・追加する', 'bold'),
        ('1. 問題のタブ（vocab、m-seifu など）で、セルを直したり行を追加したりします。', None),
        ('   行を追加するときは、似ている行をコピーして下に貼り付け、中身を書きかえると簡単です。', None),
        ('2. 最後に「manifest」タブで、直したタブの version を 1 増やします（黄色の列）。', None),
        ('   これが「新しい問題がありますよ」の合図です。増やさないとアプリは取り込みません。', None),
        ('3. アプリは次に開いたとき（インターネット接続時）に新しい問題を取り込みます。', None),
        ('   すぐ確認したいとき：KIDS「せってい → もんだいデータを こうしん」／TEEN「⚙ → 問題データを更新」', None),
        ('', None),
        ('■ 注意', 'bold'),
        ('・1行目（列の名前）とタブの名前は変えないでください。', None),
        ('・2行目（灰色）は列の説明です。アプリは読みません。', None),
        ('・_path 列が空の行はアプリが無視します（メモに使えます）。', None),
        ('・空のセルは「なし」の意味です。', None),
        ('・新しい絵や音声はここだけでは追加できません（アプリの更新が必要です）。', None),
        ('・まちがいがあってもアプリはこわれず、前の問題のまま動きます。', None),
        ('', None),
        ('■ 課金（有料・無料）の設定 → pricing タブ', 'bold'),
        ('・enabled を true にすると課金が有効になり、無料の範囲以外にカギがつきます。false のあいだは全部無料です。', None),
        ('・defaultFreeCount = 各ゲーム（ステージ）で無料で遊べるモード（レベル）の数。', None),
        ('・pricing タブは version を上げなくても、アプリの次回起動時に反映されます。', None),
        ('', None),
        ('■ _path（グループ）の意味', 'bold'),
    ]
    for app, pack, title, paths in path_notes:
        lines.append((f'{app} / {pack}（{title}）: ' + '、'.join(f'{p} = {describe_path(p)}' for p in paths), None))
    for r, (text, style) in enumerate(lines, 1):
        c = guide.cell(row=r, column=1, value=text)
        if style == 'bold':
            c.font = Font(bold=True, size=12 if r > 1 else 14)
    guide.column_dimensions['A'].width = 120

    wb.save(out)
    print(f'{out}: {len(mani_rows) - 1} packs, {sum(len(read_csv(p)) - 1 for a, root in apps for p in (Path(root) / "tools" / "sheets").glob("*.csv") if p.stem != "manifest")} questions')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
