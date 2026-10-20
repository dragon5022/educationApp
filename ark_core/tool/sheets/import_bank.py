"""Imports the school's question bank workbook (ARK_BUSTERS_実装用データ) and
voice package into the apps' question packs.

    python tool/sheets/import_bank.py <bank.xlsx> <voice_root> <kids_app> <teen_app> [--paid-out DIR]

- KIDS: vocab / listening / sentence / talk packs are rebuilt from the
  workbook (every question carries `free: true/false`, voice file names
  `KIDS_V_0001.mp3` …). Voice clips of free questions are copied into the
  app's assets; clips of paid questions go to --paid-out (default
  <kids_app>/../content/kids_voice/paid) for upload to Firebase Storage
  (`media/audio/voice/...`), from where the app downloads them after the
  purchase. Existing vocabulary pictures are kept by matching the word.
- TEEN: the 21 stage packs get their `levels` replaced (settings kept).

Pack settings (timers, counts, titles, images) come from the current pack
files, so run this on a checkout that already has them.
"""
import json
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

GRADE = {'5級': 'g5', '4級': 'g4', '3級': 'g3'}
GRADE_NUM = {'5級': 5, '4級': 4, '3級': 3}
VOICE_DIR = {'V': 'voca', 'L': 'lis', 'S': 'gram', 'T': 'talk'}
# TEEN コーナー名 → pack id
TEEN_PACKS = {
    '正負の境界線': 'm-seifu', '係数の暗号室': 'm-keisu', '指数の増幅炉': 'm-shisu',
    '図形の幾何回廊': 'm-zukei', '関数の座標宮': 'm-kansu',
    '語彙の兵器庫': 'e-goi', '文法の要塞': 'e-bunpou', '長文の戦場': 'e-choubun', '表現の司令室': 'e-hyougen',
    '漢字の間': 'k-kanji', '言葉の蔵': 'k-kotoba', '文法の塔': 'k-bunpou', '古典の庭': 'k-koten',
    '元素の実験室': 'r-genso', '生命の培養槽': 'r-inochi', '大地の記憶庫': 'r-daichi', '力学の歯車塔': 'r-rikigaku',
    '地図の迷宮': 's-chizu', '歴史の回廊': 's-rekishi', '時代の玉座': 's-jidai', '正義の評議院': 's-seigi',
}
OLD_CLIP = re.compile(r'^(\d+|g\d_\d_\d+(_r)?)\.mp3$')


def rows_of(ws):
    rows = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, '') for c in r)]
    header = [str(c or '').strip() for c in rows[2]]
    out = []
    for r in rows[3:]:
        d = {h: ('' if v is None else str(v).strip()) for h, v in zip(header, r) if h}
        if d.get('通しNo.'):
            out.append(d)
    return out


def is_free(d):
    return d.get('区分', '') != '有料'


def voice_path(d):
    name = d.get('音声ファイル名', '')
    m = re.match(r'KIDS_([VLST])_\d{4}\.mp3$', name)
    if not m:
        raise SystemExit(f'unexpected voice file name {name!r} (No.{d.get("通しNo.")})')
    return f'audio/voice/{VOICE_DIR[m.group(1)]}/{name}'


def emoji_only(s):
    return bool(s) and not re.search(r'[A-Za-z0-9぀-ヿ一-鿿]', s)


def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def save(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    bank, voice_root, kids, teen = (Path(a) for a in args[:4])
    paid_out = Path(sys.argv[sys.argv.index('--paid-out') + 1]) if '--paid-out' in sys.argv else kids.parent / 'content' / 'kids_voice' / 'paid'
    wb = load_workbook(bank, read_only=True, data_only=True)
    ws = {w.title: w for w in wb.worksheets}
    data = kids / 'assets' / 'data'
    used_clips = {}  # relative path → (src file, free)

    def clip(d):
        rel = voice_path(d)
        src = voice_root / d['音声フォルダ'] / d['音声ファイル名']
        if not src.exists():
            raise SystemExit(f'missing clip {src}')
        used_clips[rel] = (src, is_free(d))
        return rel

    # ---------------- KIDS vocab
    old = load(data / 'vocab.json')
    img_by_word = {i['w'].strip().lower(): i['img'] for i in old['items'] if i.get('img')}
    items = []
    for d in rows_of(ws['① KIDS ボキャブラリー']):
        w = d['英語'].strip()
        items.append({
            'id': int(d['通しNo.']), 'w': w, 'ja': d['日本語・意味'].strip(),
            'lv': GRADE_NUM[d['級']], 'img': img_by_word.get(w.lower()),
            'voice': clip(d), 'free': is_free(d),
        })
    old['items'] = items
    old['version'] = old.get('version', 1) + 1
    save(data / 'vocab.json', old)
    print(f'vocab: {len(items)} ({sum(i["free"] for i in items)} free, {sum(1 for i in items if i["img"])} with pictures)')

    # ---------------- KIDS listening
    old = load(data / 'listening.json')
    grades = defaultdict(lambda: defaultdict(list))
    for d in rows_of(ws['② KIDS リスニング']):
        g = GRADE[d['級']]; lv = str(int(float(d['レベル'])))
        ans = d['正解（絵）'].strip()
        opts = [o.strip() for o in re.split(r'\s*/\s*', d['まちがい選択肢']) if o.strip()]
        n = len(grades[g][lv])
        grades[g][lv].append({
            'id': f'{g}_{lv}_{n}', 't': 'pic' if emoji_only(ans) else 'text',
            'en': d['英語'].strip(), 'ja': d['日本語・意味'].strip(), 'tip': d.get('ヒント', '').strip(),
            'ans': ans, 'opts': opts, 'voice': clip(d), 'free': is_free(d),
        })
    old['grades'] = {g: {'levels': {lv: grades[g][lv] for lv in sorted(grades[g])}} for g in ['g5', 'g4', 'g3']}
    old['version'] = old.get('version', 1) + 1
    save(data / 'listening.json', old)
    n = sum(len(v) for g in grades.values() for v in g.values())
    print(f'listening: {n} ({sum(i["free"] for g in grades.values() for v in g.values() for i in v)} free)')

    # ---------------- KIDS sentence
    old = load(data / 'sentence.json')
    by_grade = defaultdict(list)
    for d in rows_of(ws['③ KIDS センテンス']):
        g = GRADE[d['級']]
        en = d['英語'].strip()
        words = en.split()
        n = len(by_grade[g])
        by_grade[g].append({
            'id': f'{g}_{n}',
            'ja': [{'t': d['日本語・意味'].strip(), 'pick': False}],
            'answer': [{'w': w} for w in words],
            'tip': d.get('ヒント', '').strip(), 'voice': clip(d), 'free': is_free(d),
        })
    old['grades'] = {g: by_grade[g] for g in ['g5', 'g4', 'g3']}
    old['version'] = old.get('version', 1) + 1
    save(data / 'sentence.json', old)
    print(f'sentence: {sum(len(v) for v in by_grade.values())} ({sum(i["free"] for v in by_grade.values() for i in v)} free)')

    # ---------------- KIDS talk
    old = load(data / 'talk.json')
    by_grade = defaultdict(list)
    for d in rows_of(ws['④ KIDS トーク']):
        g = GRADE[d['級']]
        n = len(by_grade[g])
        it = {'id': f'{g}_{n}', 'ja': d['日本語・意味'].strip(), 'en': d['英語'].strip(),
              'voice': clip(d), 'free': is_free(d)}
        if d.get('場面', '').strip():
            it['scene'] = d['場面'].strip()
        by_grade[g].append(it)
    old['grades'] = {g: by_grade[g] for g in ['g5', 'g4', 'g3']}
    old['version'] = old.get('version', 1) + 1
    save(data / 'talk.json', old)
    print(f'talk: {sum(len(v) for v in by_grade.values())} ({sum(i["free"] for v in by_grade.values() for i in v)} free)')

    # ---------------- voice clips
    free_n = paid_n = removed = 0
    for sub in VOICE_DIR.values():
        folder = kids / 'assets' / 'audio' / 'voice' / sub
        for f in folder.glob('*.mp3'):
            if OLD_CLIP.match(f.name) or (f.name.startswith('KIDS_') and f'audio/voice/{sub}/{f.name}' not in used_clips):
                f.unlink(); removed += 1
    for rel, (src, free) in used_clips.items():
        dst = (kids / 'assets' / rel) if free else (paid_out / rel)
        dst.parent.mkdir(parents=True, exist_ok=True)
        if not dst.exists() or dst.stat().st_size != src.stat().st_size:
            shutil.copyfile(src, dst)
        if free: free_n += 1
        else: paid_n += 1
    print(f'clips: {free_n} bundled (free), {paid_n} staged for Storage in {paid_out}, {removed} old clips removed')

    # ---------------- TEEN
    teen_data = teen / 'assets' / 'data'
    per_pack = defaultdict(lambda: defaultdict(list))
    for title in ['⑤ TEEN 数学', '⑥ TEEN 英語', '⑦ TEEN 国語', '⑧ TEEN 理科', '⑨ TEEN 社会']:
        for d in rows_of(ws[title]):
            pid = TEEN_PACKS[d['コーナー名']]
            choices = [d.get(f'選択肢{i}', '').strip() for i in range(1, 5)]
            choices = [c for c in choices if c]
            per_pack[pid][d['レベル']].append({
                'q': d['問題文'].strip(), 'a': d['正解'].strip(), 'choices': choices,
                'e': d.get('解説', '').strip(), 'free': is_free(d),
            })
    total = 0
    for pid, levels in per_pack.items():
        pack = load(teen_data / f'{pid}.json')
        pack['levels'] = {lv: levels[lv] for lv in sorted(levels)}
        pack['version'] = pack.get('version', 1) + 1
        save(teen_data / f'{pid}.json', pack)
        n = sum(len(v) for v in levels.values()); total += n
        print(f'  {pid}: {n} ({sum(i["free"] for v in levels.values() for i in v)} free) levels={list(levels)}')
    print(f'teen: {total} questions in {len(per_pack)} packs')
    missing = set(TEEN_PACKS.values()) - set(per_pack)
    if missing:
        print('WARNING packs without rows:', missing)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
